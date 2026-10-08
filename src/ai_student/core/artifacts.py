"""作业产出：把 AI 完成的答案写成文件，放到用户指定的文件夹。

支持 md / txt / docx；docx 由 python-docx 生成（无需 Office）。
"""
from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path

from ..config.models import OutputConfig, PersonalInfo
from ..config.paths import sanitize_filename
from .models import Artifact, ClassTask, SessionInfo

log = logging.getLogger(__name__)

_STATUS_CN = {
    "discovered": "已发现",
    "queued": "排队中",
    "working": "作答中",
    "solved": "已完成",
    "failed": "失败",
    "skipped": "已跳过",
}


class ArtifactWriter:
    """按 OutputConfig 把任务答案落盘。"""

    def __init__(self, output_cfg: OutputConfig, person: PersonalInfo) -> None:
        self.cfg = output_cfg
        self.person = person

    # ------------------------------------------------------------ 目录
    def base_dir(self) -> Path:
        d = Path(self.cfg.output_dir) if self.cfg.output_dir else Path.cwd() / "AI学生作业"
        return d

    def session_dir(self, info: SessionInfo) -> Path:
        base = self.base_dir()
        if not self.cfg.create_session_subfolder:
            return base
        stamp = (info.started_at or "")[:10] or datetime.now().strftime("%Y-%m-%d")
        subject = info.subject or "课堂"
        name = self._render(self.cfg.naming_template, info, None) or f"{stamp}_{subject}"
        return base / sanitize_filename(name, "课堂", 60)

    def _render(self, template: str, info: SessionInfo, task: ClassTask | None) -> str:
        started = info.started_at or datetime.now().isoformat()
        date = started[:10]
        time_part = started[11:16].replace(":", "") if len(started) > 15 else ""
        subject = info.subject or self.person.subject or "课堂"
        mapping = {
            "date": date,
            "time": time_part,
            "subject": subject,
            "session": info.title or info.id,
            "title": info.title or subject,
            "teacher": info.teacher or self.person.teacher_name,
            "student": self.person.name or "学生",
            "task": task.title if task else "",
        }
        out = template or "{date}_{subject}"
        for key, val in mapping.items():
            out = out.replace("{" + key + "}", str(val))
        return out.strip(" _-")

    # ------------------------------------------------------------ 单任务
    def write_task(
        self,
        info: SessionInfo,
        task: ClassTask,
        answer: str,
        workpaper: str = "",
        file_title: str = "",
        index: int = 1,
    ) -> list[Artifact]:
        folder = self.session_dir(info)
        folder.mkdir(parents=True, exist_ok=True)
        title = sanitize_filename(file_title or task.title or f"任务{index}", f"任务{index}", 60)
        stem = f"{index:02d}_{title}"
        artifacts: list[Artifact] = []

        body_parts = [answer.strip()]
        if self.cfg.include_workpaper and workpaper.strip():
            body_parts.append("\n---\n\n## 解题过程\n\n" + workpaper.strip())

        for fmt in self._formats():
            path = folder / f"{stem}.{fmt}"
            try:
                if fmt == "md":
                    self._write_md(path, info, task, "\n\n".join(body_parts))
                elif fmt == "txt":
                    self._write_txt(path, info, task, "\n\n".join(body_parts))
                elif fmt == "docx":
                    self._write_docx(path, info, task, "\n\n".join(body_parts))
                else:
                    continue
                artifacts.append(
                    Artifact(
                        task_id=task.id, path=str(path), fmt=fmt, title=task.title,
                        size_bytes=path.stat().st_size,
                    )
                )
            except PermissionError as exc:
                log.error("写入 %s 被拒绝（文件可能正被 Word 打开）: %s", path, exc)
            except Exception as exc:
                log.exception("写入 %s 失败: %s", path, exc)
        return artifacts

    def _formats(self) -> list[str]:
        allowed = {"md", "txt", "docx"}
        fmts = [f.lower().lstrip(".") for f in (self.cfg.write_formats or ["md"])]
        out = [f for f in fmts if f in allowed]
        return out or ["md"]

    # ------------------------------------------------------------ 头部信息
    def _meta_lines(self, info: SessionInfo, task: ClassTask) -> list[str]:
        started = (info.started_at or "").replace("T", " ")[:16]
        lines = []
        header = self.person.header_line()
        if header:
            lines.append(header)
        row2 = []
        if info.subject or self.person.subject:
            row2.append(f"科目：{info.subject or self.person.subject}")
        if info.teacher or self.person.teacher_name:
            row2.append(f"任课教师：{info.teacher or self.person.teacher_name}")
        if started:
            row2.append(f"上课时间：{started}")
        if row2:
            lines.append("　".join(row2))
        lines.append(f"任务类型：{task.task_type}" + (f"　截止：{task.deadline}" if task.deadline else ""))
        if task.requirements:
            lines.append("要求：" + "；".join(r for r in task.requirements if r))
        return [x for x in lines if x.strip()]

    def _write_md(self, path: Path, info: SessionInfo, task: ClassTask, body: str) -> None:
        parts = [f"# {task.title}", ""]
        parts.extend(f"> {line}" for line in self._meta_lines(info, task))
        parts.append("")
        parts.append(body.strip())
        if self.person.sign_artifacts and self.person.name:
            parts.append("")
            parts.append(f"> 完成人：{self.person.name}　生成时间：{datetime.now():%Y-%m-%d %H:%M}")
        path.write_text("\n".join(parts) + "\n", encoding="utf-8")

    def _write_txt(self, path: Path, info: SessionInfo, task: ClassTask, body: str) -> None:
        lines = [task.title, "=" * min(40, max(8, len(task.title) * 2)), ""]
        lines.extend(self._meta_lines(info, task))
        lines.append("")
        lines.append(_md_to_plain(body))
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    def _write_docx(self, path: Path, info: SessionInfo, task: ClassTask, body: str) -> None:
        from docx import Document  # 延迟导入，加快启动

        doc = Document()
        doc.add_heading(task.title or "作业", level=1)
        for line in self._meta_lines(info, task):
            p = doc.add_paragraph(line)
            p.runs[0].italic = True if p.runs else False
        doc.add_paragraph("")
        _md_into_docx(doc, body)
        if self.person.sign_artifacts and self.person.name:
            doc.add_paragraph("")
            doc.add_paragraph(f"完成人：{self.person.name}　生成时间：{datetime.now():%Y-%m-%d %H:%M}")
        doc.save(str(path))

    # ------------------------------------------------------------ 汇总
    def write_summary(
        self,
        info: SessionInfo,
        tasks: list[ClassTask],
        notes_markdown: str = "",
        extra_transcript: str = "",
    ) -> Artifact | None:
        folder = self.session_dir(info)
        folder.mkdir(parents=True, exist_ok=True)
        stamp = (info.started_at or "")[:10] or datetime.now().strftime("%Y-%m-%d")
        path = folder / f"00_{stamp}_作业汇总.md"
        lines = [f"# {stamp} 课堂作业汇总", ""]
        header = self.person.header_line()
        if header:
            lines.append(f"> {header}")
        lines.append(f"> 会话：{info.display_name}")
        lines.append("")
        lines.append(f"共发现 {len(tasks)} 项任务：")
        lines.append("")
        lines.append("| # | 任务 | 类型 | 状态 | 产出文件 |")
        lines.append("|---|------|------|------|----------|")
        for i, t in enumerate(tasks, 1):
            files_txt = "<br>".join(Path(a).name for a in t.artifacts) or "-"
            title = (t.title or "").replace("|", "/")
            lines.append(
                f"| {i} | {title} | {t.task_type} | {_STATUS_CN.get(t.status, t.status)} | {files_txt} |"
            )
        lines.append("")
        pending = [t for t in tasks if t.status not in ("solved",)]
        if pending:
            lines.append("## 待处理")
            lines.append("")
            for t in pending:
                lines.append(f"- **{t.title}**（{_STATUS_CN.get(t.status, t.status)}）{t.detail[:120]}")
            lines.append("")
        if self.cfg.include_transcript and notes_markdown:
            lines.append("---")
            lines.append("")
            lines.append(notes_markdown.strip())
        if extra_transcript:
            lines.append("")
            lines.append("---")
            lines.append("")
            lines.append(extra_transcript.strip())
        try:
            path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        except Exception as exc:
            log.error("写汇总失败: %s", exc)
            return None
        return Artifact(task_id="", path=str(path), fmt="md", title="作业汇总",
                        size_bytes=path.stat().st_size)

    def write_manual_notes(self, info: SessionInfo, notes_markdown: str, transcript_md: str = "") -> list[str]:
        """把课堂笔记/时间轴也放进输出目录（便于用户查阅）。"""
        if not self.cfg.include_transcript:
            return []
        folder = self.session_dir(info)
        folder.mkdir(parents=True, exist_ok=True)
        out: list[str] = []
        if notes_markdown.strip():
            p = folder / "01_课堂笔记.md"
            p.write_text(notes_markdown, encoding="utf-8")
            out.append(str(p))
        if transcript_md.strip():
            p = folder / "02_课堂时间轴.md"
            p.write_text(transcript_md, encoding="utf-8")
            out.append(str(p))
        return out


# ------------------------------------------------------------------ Markdown 工具
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_INLINE_CODE_RE = re.compile(r"`([^`]+)`")
_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _md_to_plain(md: str) -> str:
    out_lines: list[str] = []
    in_code = False
    for line in (md or "").splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            out_lines.append("    " + line)
            continue
        s = line
        s = re.sub(r"^#{1,6}\s*", "", s)
        s = _BOLD_RE.sub(r"\1", s)
        s = _INLINE_CODE_RE.sub(r"\1", s)
        s = _LINK_RE.sub(r"\1（\2）", s)
        out_lines.append(s)
    return "\n".join(out_lines)


def _md_into_docx(doc, md: str) -> None:
    """极简 Markdown → docx：标题/列表/代码块/引用/表格分隔线。"""
    in_code = False
    code_buf: list[str] = []

    def flush_code() -> None:
        nonlocal code_buf
        if code_buf:
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_buf))
            run.font.name = "Consolas"
            code_buf = []

    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_buf.append(line)
            continue
        if not line.strip():
            doc.add_paragraph("")
            continue
        if line.strip() in ("---", "***", "___"):
            doc.add_paragraph("——————")
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line.strip())
        if m:
            doc.add_heading(_strip_inline(m.group(2)), level=min(4, len(m.group(1))))
            continue
        if line.lstrip().startswith(("- ", "* ", "+ ")):
            doc.add_paragraph(_strip_inline(line.lstrip()[2:]), style="List Bullet")
            continue
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if m:
            doc.add_paragraph(_strip_inline(m.group(2)), style="List Number")
            continue
        if line.lstrip().startswith(">"):
            p = doc.add_paragraph(_strip_inline(line.lstrip()[1:].strip()))
            for r in p.runs:
                r.italic = True
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            doc.add_paragraph("　".join(_strip_inline(c) for c in cells))
            continue
        # 普通段落：保留 **加粗**
        p = doc.add_paragraph()
        pos = 0
        for m2 in _BOLD_RE.finditer(line):
            if m2.start() > pos:
                p.add_run(_strip_inline(line[pos:m2.start()]))
            r = p.add_run(_strip_inline(m2.group(1)))
            r.bold = True
            pos = m2.end()
        if pos < len(line):
            p.add_run(_strip_inline(line[pos:]))
    flush_all = in_code
    if flush_all:
        flush_code()


def _strip_inline(text: str) -> str:
    s = _INLINE_CODE_RE.sub(r"\1", text)
    s = _LINK_RE.sub(r"\1（\2）", s)
    s = _BOLD_RE.sub(r"\1", s)
    return s.replace("**", "")
